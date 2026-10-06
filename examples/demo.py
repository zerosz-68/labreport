#!/usr/bin/env python
"""端到端演示：现场造一份"带合并单元格的模板 + 实验数据"，把整条流水线跑一遍。

不需要任何真实实验报告，也不会联网。跑法：

    python examples/demo.py            # 全部步骤
    python examples/demo.py --keep     # 保留 examples/_out/（默认保留，便于打开看）

演示内容：
    造模板 -> inspect 看物理网格 -> 造 data.json -> data check 校验
    -> fill 先出计划再落地 -> audit 自检
"""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "src"))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from labreport.cli import main as cli  # noqa: E402

OUT = os.path.join(HERE, "_out")
TPL = os.path.join(OUT, "实验报告模板.docx")
DATA = os.path.join(OUT, "data.json")
MAP = os.path.join(OUT, "map.json")
FILLED = os.path.join(OUT, "成品.docx")


def step(n, total, title):
    print()
    print(f"[{n}/{total}] {title}")
    print("-" * 64)


def pretty(path: str) -> str:
    """把工作目录里的绝对路径缩成相对路径，命令回显更好读。"""
    try:
        rel = os.path.relpath(path, ROOT)
    except ValueError:
        return path
    return rel if not rel.startswith("..") else path


def call(argv, show=6, tail=False):
    """在进程内跑一条 CLI 命令，打印它的部分输出与退出码。"""
    print("    $ labreport " + " ".join(pretty(str(a)) for a in argv))
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = cli(argv)
    lines = [l for l in buf.getvalue().splitlines() if l.strip()]
    picked = lines[-show:] if tail else lines[:show]
    for line in picked:
        # 仅为了演示好读：把长绝对路径缩写（内容不变）
        print("      " + line.replace(OUT + os.sep, "examples/_out/").replace(OUT, "examples/_out"))
    if len(lines) > show:
        print(f"      …（共 {len(lines)} 行）")
    print(f"      退出码 = {code}")
    return code


def make_template():
    """造一份带水平合并单元格的模板：信息表 + 数据表 + 公式区。"""
    doc = Document()
    doc.add_heading("实验实训报告（演示用）", level=0)
    doc.add_paragraph("本文件由 examples/demo.py 现场生成，用于演示 labreport 的填写与自检。")

    info = doc.add_table(rows=2, cols=4)
    info.style = "Table Grid"
    info.cell(0, 1).merge(info.cell(0, 2))  # 水平合并：实验名称跨两列
    info.cell(1, 1).merge(info.cell(1, 2))  # 水平合并：班级跨两列
    for r, c, text in [
        (0, 0, "实验名称"), (0, 3, "成绩"),
        (1, 0, "班级"), (1, 3, "姓名"),
    ]:
        info.cell(r, c).text = text

    data = doc.add_table(rows=4, cols=5)
    data.style = "Table Grid"
    header = ["试样", "质量 m/g", "尺寸 D/mm", "周期 T/ms", "J理/(kg·m²)"]
    for c, text in enumerate(header):
        data.cell(0, c).text = text
    data.cell(1, 0).text = "圆柱"
    for r in (2, 3):  # 后两行留成整行合并的公式区
        data.cell(r, 0).merge(data.cell(r, 4))

    core = doc.styles["Normal"].font
    core.name = "宋体"
    core.size = Pt(10.5)
    doc.save(TPL)
    print(f"      已生成: {os.path.relpath(TPL, ROOT)}")


def make_data():
    data = {
        "experiment": "扭摆法测转动惯量（演示数据）",
        "source": {"images": ["（演示，无照片）"], "note": "数值为演示用，非真实测量"},
        "quantities": {
            "m_cyl": {"value": 100.00, "unit": "g", "source": "演示"},
            "D_cyl": {"value": 50.00, "unit": "mm", "source": "演示"},
            "K": {"value": 0.1234, "unit": "N*m/rad", "source": "演示"},
        },
        "measurements": [
            {
                "id": "T1",
                "label": "空盘+圆柱",
                "unit": "ms",
                "values": [100.0, 100.1, 99.9, 100.0, 100.1, 99.9],
                "mean": 100.0,
                "precision": 1,
                "source": "演示",
            }
        ],
        "targets": [
            {"name": "J1理", "expr": "1/8*m_cyl*D_cyl^2", "unit": "kg*m^2",
             "expected": 3.125e-5, "tolerance": 0.01},
            {"name": "J1实", "expr": "K*T1^2/(4*pi^2)", "unit": "kg*m^2",
             "expected": 3.1257e-5, "tolerance": 0.01},
            {"name": "E", "expr": "abs(J1实-J1理)/J1理*100", "unit": "%",
             "expected": 0.0243, "tolerance": 0.5},
        ],
    }
    with open(DATA, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)

    mapping = {
        "_说明": "table/row/col 是屏幕上看到的物理网格坐标（1 起）",
        "data": os.path.basename(DATA),
        "targets": [
            {"table": 1, "row": 1, "col": 2, "set": "{{experiment}}"},
            {"table": 1, "row": 2, "col": 2, "set": "演示班"},
            {"table": 2, "row": 2, "col": 2, "set": "{{quantities.m_cyl}}", "format": "{:.2f}"},
            {"table": 2, "row": 2, "col": 3, "set": "{{quantities.D_cyl}}", "format": "{:.2f}"},
            {"table": 2, "row": 2, "col": 4, "set": "{{measurements.T1}}", "format": "{:.1f}"},
            {"table": 2, "row": 2, "col": 5, "set": "{{targets.J1理}}", "format": "{:.4e}"},
            {"table": 2, "row": 3, "col": 1,
             "append": "公式（仪器常数 K = {{quantities.K}} N·m/rad）："},
            {"table": 2, "row": 3, "col": 1, "append": True,
             "formula": "J_{1\\mathrm{理}}=\\frac{1}{8}mD^{2}={{targets.J1理}}"},
            {"table": 2, "row": 4, "col": 1, "append": True,
             "formula": "J_{1\\mathrm{实}}=\\frac{KT^{2}}{4\\pi^{2}}={{targets.J1实}}"},
            {"table": 2, "row": 4, "col": 1, "append": True,
             "formula": "E=\\frac{\\left|J_{1\\mathrm{实}}-J_{1\\mathrm{理}}\\right|}{J_{1\\mathrm{理}}}\\times 100\\%={{targets.E}}"},
        ],
    }
    with open(MAP, "w", encoding="utf-8") as fh:
        json.dump(mapping, fh, ensure_ascii=False, indent=2)
    print(f"      已生成: {os.path.relpath(DATA, ROOT)} 与 {os.path.relpath(MAP, ROOT)}")


def main():
    ap = argparse.ArgumentParser(description="labreport 端到端演示")
    ap.add_argument("--keep", action="store_true", help="保留输出目录（默认保留）")
    ap.parse_args()

    total = 6
    os.makedirs(OUT, exist_ok=True)
    print("labreport 端到端演示")
    print(f"输出目录: {os.path.relpath(OUT, ROOT)}")

    step(1, total, "造一份带合并单元格的模板")
    make_template()

    step(2, total, "inspect：看清表格物理网格（合并单元格会被标出来）")
    call(["inspect", TPL, "--max-text", "12", "--formula-sample", "0"], show=14)

    step(3, total, "造数据与填写映射（真实场景里由 agent 读照片后生成）")
    make_data()

    step(4, total, "data check：单位归一化 / 平均值自洽 / 公式重算比对")
    call(["data", "check", DATA], show=16)

    step(5, total, "fill：先看计划（不写文件），确认后再落地")
    call(["fill", TPL, "--map", MAP, "--data", DATA, "-o", FILLED], show=8)
    call(["fill", TPL, "--map", MAP, "--data", DATA, "-o", FILLED, "--auto"], tail=True, show=3)

    step(6, total, "audit：交付前自检（结构 / 公式 / 占位 / 数据回读）")
    call(["audit", FILLED, "--source", TPL, "--data", DATA], show=10)

    print()
    print("=" * 64)
    print("演示结束。产物：")
    for p in (TPL, DATA, MAP, FILLED):
        print(f"  {os.path.relpath(p, ROOT)}")
    print()
    print("接着可以自己看排版（需要 Word 或 LibreOffice）：")
    print(f"  labreport pdf \"{os.path.relpath(FILLED, ROOT)}\" --pages")
    print("然后把你自己的报告套进来：docs/quickstart.md")


if __name__ == "__main__":
    main()
