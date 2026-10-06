"""成品自检：结构有没有被破坏、公式是不是原生、占位有没有残留、数据填全了没有。

定位：`fill` 之后的自检门禁（也是交付前的最后一道关）。回答四个问题：
1. 与 source 模板相比，表格几何/结构有没有被破坏？
2. 公式是不是真的 Word 原生（OMML），有没有 LaTeX 残留或空公式？
3. 还有没有没填的占位（`{{...}}`、`—`、`（ ）`、`____`）？
4. data.json 里的每个数值，在成品里能不能找到（漏填/写错形式的检查）？
"""

from __future__ import annotations

import json
import math
import os
import re

from docx import Document

from ..data.check import check as check_data
from ..docx.inspect import M, W, table_structure

# 报告里的科学计数法：1.4091×10−3 / 1.4091x10^-3 / 1.4091*10-3
SCI_RE = re.compile(r"(\d+(?:\.\d+)?)\s*[×xX*]\s*10\s*\^?\s*([−\-+]?)\s*(\d+)")
NUM_RE = re.compile(r"[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?")
LATEX_RESIDUE = re.compile(r"\\(frac|sqrt|times|cdot|pi|theta|alpha|omega|mathrm|sum|int|partial)\b|\^\{|_\{")
PLACEHOLDER_RESIDUE = [
    (re.compile(r"\{\{[^}]*\}\}"), "未解析的映射占位符 {{...}}"),
    (re.compile(r"^[\s\u3000]*$"), "空单元格"),
    (re.compile(r"^[—\-–_…\.·\s]{2,}$"), "未填占位（— / ____ 等）"),
    (re.compile(r"^[（(]\s*[)）]$"), "未填占位（空括号）"),
    (re.compile(r"[【\[［]\s*[】\]］]"), "未填占位（空方括号）"),
]


def _top_tables(doc):
    return [ch for ch in doc.element.body if ch.tag == W + "tbl"]


def _text_variants(doc):
    """两种拼接方式各取一份，避免"相邻单元格数字粘连"与"数字被 run 拆开"两种坑。

    不带分隔符：能识别被拆成多个 run 的一个数（如 "1116" + ".12"）。
    带分隔符：能避免相邻单元格粘连（如 774 与 1476 变成 7741476）。
    回读时两份都试，命中任一即可。
    """
    no_sep, with_sep = [], []
    for el in doc.element.body.iter():
        if el.tag in (W + "t", M + "t") and el.text:
            no_sep.append(el.text)
            with_sep.append(el.text)
            with_sep.append(" ")
    return "".join(no_sep), "".join(with_sep)


def _all_text(doc) -> str:
    """按文档顺序拼接所有文字（w:t 与公式 m:t），供 LaTeX 残留检查。"""
    parts = []
    for el in doc.element.body.iter():
        if el.tag in (W + "t", M + "t") and el.text:
            parts.append(el.text)
    return "".join(parts)


def _numbers(text: str):
    norm = SCI_RE.sub(
        lambda m: f"{m.group(1)}e{'-' if m.group(2) in ('−', '-') else ''}{m.group(3)}",
        text,
    )
    out = []
    for tok in NUM_RE.findall(norm):
        try:
            out.append(float(tok))
        except ValueError:
            pass
    return out


def _close(a: float, b: float, tol: float) -> bool:
    if b == 0:
        return abs(a) <= tol
    return abs(a - b) / abs(b) <= tol


def structure(path: str) -> dict:
    doc = Document(path)
    tables = [table_structure(t) for t in _top_tables(doc)]
    return {
        "paragraphs": len(doc.paragraphs),
        "nonEmptyParagraphs": len([p for p in doc.paragraphs if p.text.strip()]),
        "sections": len(doc.sections),
        "tables": [
            {
                "rows": t["rows"],
                "cols": t["gridCols"],
                "merged": t["mergedCells"],
                "nested": t["nestedTables"],
                "blank": t["blankCells"],
            }
            for t in tables
        ],
        "mathCount": sum(1 for _ in doc.element.body.iter(M + "oMath")),
        "images": sum(1 for _ in doc.element.body.iter(W + "drawing")),
    }


def _variants(item: str) -> list:
    """提纲条目的宽容匹配形式（成品里可能省掉“实验”二字）。"""
    outs = [item]
    if item.startswith("实验") and len(item) > 2:
        outs.append(item[2:])
    outs.append(item.replace("实验", ""))
    return [x for x in dict.fromkeys(outs) if len(x) >= 2]


def _outline_items(source: str) -> list:
    """从模板里读出“实验报告应包括以下内容：…”这段提纲，逐条返回。

    这就是需求里说的“内容的目录在实验报告（未填写的）”——与其让使用者手抄一份
    章节清单，不如直接从模板里读，成品逐项核对。
    """
    doc = Document(source)
    for tbl in [ch for ch in doc.element.body if ch.tag == W + "tbl"]:
        for tc in tbl.iter(W + "tc"):
            paras = [
                "".join(t.text or "" for t in p.iter(W + "t")).strip()
                for p in tc.findall(W + "p")
            ]
            idx = next((i for i, s in enumerate(paras) if "应包括以下内容" in s), None)
            if idx is None:
                continue
            items = []
            head = paras[idx].split("：", 1)
            if len(head) == 2 and head[1].strip():
                items.append(head[1].strip())
            items += [s for s in paras[idx + 1:] if s]
            clean = []
            for s in items:
                s = re.sub(r"^[\s\d一二三四五六七八九十]+[.、)）]?\s*", "", s).strip()
                if s and not s.endswith("：") and s not in clean:
                    clean.append(s)
            if clean:
                return clean
    return []


def _search_text(doc) -> str:
    """正文全文，但**跳过"实验报告应包括以下内容"那个单元格**。

    否则提纲本身（实验目的、实验基本原理…）会被当成"章节已经写了"，造成假通过。

    注意：不能先收集节点再用 id() 判断——lxml 的代理对象是临时创建的，
    两次遍历拿到的 id 不同，必须一边遍历一边跳过整棵子树。
    """
    parts = []

    def walk(el, skip: bool) -> None:
        for child in el:
            if child.tag == W + "tc":
                text = "".join(t.text or "" for t in child.iter(W + "t"))
                walk(child, skip or ("应包括以下内容" in text))
            elif child.tag == W + "t":
                if not skip:
                    parts.append(child.text or "")
            else:
                walk(child, skip)

    walk(doc.element.body, False)
    return "".join(parts)


def audit(
    target: str,
    source: str | None,
    data: dict | None,
    tol: float = 0.002,
    outline: bool = True,
    require_attachment: bool = False,
    attachment_label: str = "附件",
) -> dict:
    doc = Document(target)
    issues = []
    report = {"target": os.path.abspath(target), "source": os.path.abspath(source) if source else ""}

    # ---------- 1. 结构对比
    tgt_struct = structure(target)
    report["structure"] = tgt_struct
    if source and os.path.isfile(source):
        src_struct = structure(source)
        report["sourceStructure"] = src_struct
        if len(src_struct["tables"]) != len(tgt_struct["tables"]):
            issues.append(
                {
                    "level": "error",
                    "where": "结构",
                    "message": f"顶层表格数变了：模板 {len(src_struct['tables'])} → 成品 {len(tgt_struct['tables'])}",
                }
            )
        for i, (s, t) in enumerate(zip(src_struct["tables"], tgt_struct["tables"]), 1):
            if (s["rows"], s["cols"]) != (t["rows"], t["cols"]):
                issues.append(
                    {
                        "level": "error",
                        "where": f"表{i}",
                        "message": f"网格尺寸变了：{s['rows']}×{s['cols']} → {t['rows']}×{t['cols']}",
                    }
                )
            if s["merged"] != t["merged"]:
                issues.append(
                    {
                        "level": "warn",
                        "where": f"表{i}",
                        "message": f"合并单元格数变了：{s['merged']} → {t['merged']}",
                    }
                )
        if src_struct["sections"] != tgt_struct["sections"]:
            issues.append({"level": "warn", "where": "分节", "message": "分节数变了"})

    # ---------- 2. 公式健康
    maths = list(doc.element.body.iter(M + "oMath"))
    empty_math = 0
    for om in maths:
        txt = "".join(t.text or "" for t in om.iter(M + "t")).strip()
        if not txt:
            empty_math += 1
    report["math"] = {"count": len(maths), "empty": empty_math}
    if empty_math:
        issues.append({"level": "error", "where": "公式", "message": f"{empty_math} 个公式是空的"})

    full_text = _all_text(doc)
    # LaTeX 残留：正文里出现 \frac / ^{ / _{ 说明公式没转成原生
    residue = LATEX_RESIDUE.findall(full_text)
    if residue:
        issues.append(
            {
                "level": "error",
                "where": "公式",
                "message": f"发现 LaTeX 残留 {len(residue)} 处（公式可能没转成 Word 原生）：{residue[:4]}",
            }
        )

    # ---------- 3. 占位残留
    placeholders = []
    for ti, tbl in enumerate(_top_tables(doc), 1):
        for key, a in table_structure(tbl)["anchors"].items():
            txt = a["text"]
            if a["nNested"] or a["math"] or a["nImages"]:
                continue
            for pat, desc in PLACEHOLDER_RESIDUE:
                if txt and pat.search(txt) and desc != "空单元格":
                    placeholders.append({"where": f"表{ti} r{a['r']+1}c{a['c']+1}", "text": txt, "kind": desc})
    for p in doc.paragraphs:
        t = p.text.strip()
        if t and "{{" in t:
            placeholders.append({"where": "正文段落", "text": t[:40], "kind": "未解析的映射占位符"})
    report["placeholders"] = placeholders
    for ph in placeholders:
        issues.append({"level": "warn", "where": ph["where"], "message": f"{ph['kind']}：{ph['text']}"})

    # ---------- 4. 数据回读
    if data:
        vals = check_data(data)
        ns_text, sp_text = _text_variants(doc)
        nums = _numbers(ns_text) + _numbers(sp_text)

        items = []
        for q in vals["quantities"]:
            items.append({"kind": "quantity", "name": q["key"], "value": q["value"], "tol": tol})
        for m in vals["measurements"]:
            prec = int(m.get("precision", 0) or 0)
            abs_tol = 0.5 * (10 ** (-prec))
            t_tol = max(tol, abs_tol / abs(m["mean"]) if m["mean"] else tol)
            items.append({"kind": "measurement", "name": m["id"], "value": m["mean"], "tol": t_tol})
            if m.get("declaredMean") is not None:
                items.append(
                    {
                        "kind": "measurement(记录值)",
                        "name": m["id"],
                        "value": float(m["declaredMean"]),
                        "tol": t_tol,
                    }
                )
        for t_ in vals["targets"]:
            items.append(
                {
                    "kind": "target",
                    "name": t_["name"],
                    "value": t_["value"],
                    "tol": max(tol, float(t_.get("tolerance") or 0)),
                }
            )

        found, missing = [], []
        for it in items:
            value, t_tol = it["value"], it["tol"]
            hit = next((n for n in nums if _close(n, value, t_tol)), None)
            how = "直读"
            if hit is None and value != 0:
                # 允许"只写了有效数字"的情形（如 1.4091×10⁻³ 在文里写作 1.4091）
                exp = math.floor(math.log10(abs(value)))
                mant = value / (10 ** exp)
                hit = next((n for n in nums if _close(n, mant, t_tol)), None)
                if hit is not None:
                    how = "尾数"
            rec = {
                "kind": it["kind"],
                "name": it["name"],
                "value": value,
                "tol": t_tol,
                "matched": hit,
                "how": how if hit is not None else "",
            }
            (found if hit is not None else missing).append(rec)
        report["readback"] = {"total": len(items), "found": len(found), "missing": missing}
        for miss in missing:
            issues.append(
                {
                    "level": "warn",
                    "where": f"数据回读·{miss['kind']}",
                    "message": f"{miss['name']} = {miss['value']:.6g} 未在成品中命中（漏填或写成了别的形式）",
                }
            )

    # ---------- 5. 按模板提纲核对章节
    if outline and source and os.path.isfile(source):
        items = _outline_items(source)
        if items:
            body_text = _search_text(doc)
            missing_items = [
                it for it in items if not any(v in body_text for v in _variants(it))
            ]
            report["outline"] = {"items": items, "missing": missing_items}
            if missing_items:
                issues.append(
                    {
                        "level": "error",
                        "where": "章节",
                        "message": f"模板提纲 {len(items)} 项，成品缺 {len(missing_items)} 项：{missing_items}",
                    }
                )
    elif outline:
        report["outline"] = {"items": [], "missing": [], "note": "未提供 --source，跳过提纲核对"}

    # ---------- 6. 文末附件（如课堂报告书原件；可能在表格框里，所以要全文找）
    order = list(doc.element.body.iter(W + "p"))
    label_pos = []
    for i, p in enumerate(order):
        text = "".join(t.text or "" for t in p.iter(W + "t")).strip()
        if attachment_label and (text == attachment_label or text.startswith(attachment_label)):
            label_pos.append(i)
    label_idx = label_pos[-1] if label_pos else None
    images_after = 0
    if label_idx is not None:
        for p in order[label_idx:]:
            images_after += len(list(p.iter(W + "drawing"))) + len(list(p.iter(W + "pict")))
    report["attachment"] = {
        "label": attachment_label,
        "foundLabel": label_idx is not None,
        "imagesAfter": images_after,
    }
    if require_attachment and (label_idx is None or images_after == 0):
        issues.append(
            {
                "level": "error",
                "where": "附件",
                "message": f"文末没有附件（需要一行「{attachment_label}」+ 至少 1 张图片）",
            }
        )

    errors = [i for i in issues if i["level"] == "error"]
    warns = [i for i in issues if i["level"] == "warn"]
    report["issues"] = issues
    report["summary"] = {"error": len(errors), "warn": len(warns)}
    report["ok"] = not errors
    return report


def render(r: dict) -> str:
    out = []
    s = r["structure"]
    out.append(f"成品      : {r['target']}")
    if r["source"]:
        out.append(f"对照模板  : {r['source']}")
    out.append(
        f"结构      : 段落 {s['paragraphs']}（非空 {s['nonEmptyParagraphs']}）| 分节 {s['sections']}"
        f" | 顶层表 {len(s['tables'])} | 公式 {s['mathCount']} | 图片 {s['images']}"
    )
    if "sourceStructure" in r:
        ss = r["sourceStructure"]
        out.append(f"模板对照  : 段落 {ss['paragraphs']} | 表 {len(ss['tables'])} | 公式 {ss['mathCount']}")
        out.append(
            "  表格几何: "
            + " | ".join(
                f"表{i}: {a['rows']}×{a['cols']}→{b['rows']}×{b['cols']}"
                for i, (a, b) in enumerate(zip(ss["tables"], s["tables"]), 1)
            )
        )
    if "readback" in r:
        rb = r["readback"]
        out.append(f"数据回读  : {rb['found']}/{rb['total']} 命中")
        for m in rb["missing"]:
            out.append(f"    ✗ {m['name']} = {m['value']:.6g}")
    if r.get("outline", {}).get("items"):
        ol = r["outline"]
        out.append(
            f"章节提纲  : {len(ol['items']) - len(ol['missing'])}/{len(ol['items'])} 项齐备"
            + (f"    缺: {ol['missing']}" if ol["missing"] else "")
        )
    att = r.get("attachment")
    if att:
        out.append(
            "文末附件  : "
            + (
                f"找到「{att['label']}」，其后 {att['imagesAfter']} 张图片"
                if att["foundLabel"]
                else "未找到附件标签"
            )
        )
    out.append("")
    if r["issues"]:
        order = {"error": 0, "warn": 1}
        for i in sorted(r["issues"], key=lambda x: order.get(x["level"], 9)):
            out.append(f"  [{i['level']:<5}] {i['where']}: {i['message']}")
    else:
        out.append("  没有问题")
    out.append("")
    out.append(
        f"结论: {'可以交付' if r['ok'] else '不可交付（有 error）'}"
        f"（error {r['summary']['error']} / warn {r['summary']['warn']}）"
    )
    return "\n".join(out)


def run(args) -> int:
    if not os.path.isfile(args.target):
        print(f"找不到成品：{args.target}")
        return 1
    data = None
    if getattr(args, "data", None):
        if not os.path.isfile(args.data):
            print(f"找不到数据文件：{args.data}")
            return 1
        with open(args.data, encoding="utf-8") as fh:
            data = json.load(fh)
    source = getattr(args, "source", None)
    if source and not os.path.isfile(source):
        print(f"找不到对照模板：{source}")
        return 1

    r = audit(
        args.target,
        source,
        data,
        tol=float(getattr(args, "tolerance", 0.002)),
        outline=not getattr(args, "skip_outline", False),
        require_attachment=bool(getattr(args, "require_attachment", False)),
        attachment_label=getattr(args, "attachment_label", "附件") or "附件",
    )
    if getattr(args, "json", False):
        print(json.dumps(r, ensure_ascii=False, indent=2, default=str))
    else:
        print(render(r))
    return 0 if r["ok"] else 1
